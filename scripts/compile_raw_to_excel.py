#!/usr/bin/env python3
"""
Compile validated raw JSON/JSONL files under data/ and data/balad/ into one multi-sheet Excel workbook.

Each operational source becomes its own worksheet in data/combined_raw_data.xlsx.

Usage:
  docker compose exec -T backend python scripts/compile_raw_to_excel.py
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger("compile-raw-to-excel")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA_DIR = PROJECT_ROOT / "data"
DEFAULT_OUTPUT = DEFAULT_DATA_DIR / "combined_raw_data.xlsx"

SKIP_NAME_PATTERN = re.compile(r"test", re.IGNORECASE)
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
EXCEL_SHEET_MAX_LEN = 31
MAX_COLUMN_WIDTH = 80

def _configure_logging(verbose: bool) -> None:
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        stream=sys.stderr,
    )

def discover_data_files(data_dir: Path) -> list[Path]:
    candidates: list[Path] = []
    for scan_dir in (data_dir, data_dir / "balad"):
        if scan_dir.is_dir():
            candidates.extend(scan_dir.glob("*.json"))
            candidates.extend(scan_dir.glob("*.jsonl"))
    return sorted({path.resolve() for path in candidates})

def _has_failed_status(data: Any) -> bool:
    if not isinstance(data, dict):
        return False
    if data.get("status") == "failed":
        return True
    crawl_meta = data.get("crawl_metadata")
    return isinstance(crawl_meta, dict) and crawl_meta.get("status") == "failed"

def should_skip_file(path: Path, data: Any | None = None) -> str | None:
    if SKIP_NAME_PATTERN.search(path.stem):
        return "filename contains 'test'"
    if path.stat().st_size == 0:
        return "file is empty"
    if data is not None and _has_failed_status(data):
        return "crawler status is failed"
    return None

def make_sheet_name(path: Path, used_names: set[str]) -> str:
    stem = path.stem
    base = stem[:-4] if stem.endswith("_raw") else stem
    base = INVALID_SHEET_CHARS.sub("_", base).strip() or "sheet"
    if len(base) > EXCEL_SHEET_MAX_LEN:
        base = base[:EXCEL_SHEET_MAX_LEN]

    candidate = base
    counter = 1
    while candidate.casefold() in {name.casefold() for name in used_names}:
        suffix = f"_{counter}"
        trimmed = base[: EXCEL_SHEET_MAX_LEN - len(suffix)]
        candidate = f"{trimmed}{suffix}"
        counter += 1

    used_names.add(candidate)
    return candidate

def _flatten_record(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flat: dict[str, Any] = {}
    for key, value in record.items():
        column = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            flat.update(_flatten_record(value, column))
        elif isinstance(value, list):
            if value and all(isinstance(item, (str, int, float, bool)) or item is None for item in value):
                flat[column] = ", ".join(str(item) for item in value if item is not None)
            else:
                flat[column] = json.dumps(value, ensure_ascii=False)
        else:
            flat[column] = value
    return flat

def _format_list_cell(values: list[Any]) -> str:
    parts: list[str] = []
    for item in values:
        if isinstance(item, (dict, list)):
            parts.append(json.dumps(item, ensure_ascii=False))
        elif item is None:
            continue
        else:
            parts.append(str(item).strip())
    return ", ".join(part for part in parts if part)

def serialize_cell(value: Any) -> Any:
    if value is None:
        return value
    if isinstance(value, float) and pd.isna(value):
        return value
    if isinstance(value, list):
        return _format_list_cell(value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value

def coerce_excel_text(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for column in out.columns:
        out[column] = out[column].map(
            lambda value: ""
            if value is None or (isinstance(value, float) and pd.isna(value))
            else str(value)
        )
    return out

def autofit_columns(worksheet: Any, *, max_width: int = MAX_COLUMN_WIDTH) -> None:
    for column_index, column_cells in enumerate(
        worksheet.iter_cols(1, worksheet.max_column),
        start=1,
    ):
        max_length = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        width = min(max(max_length + 2, 10), max_width)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

def style_header_row(worksheet: Any) -> None:
    header_alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for column_index in range(1, worksheet.max_column + 1):
        worksheet.cell(row=1, column=column_index).alignment = header_alignment

def prepare_dataframe(records: list[dict[str, Any]]) -> pd.DataFrame:
    if not records:
        return pd.DataFrame()
    df = pd.DataFrame(records)
    for column in df.columns:
        df[column] = df[column].apply(serialize_cell)
    return coerce_excel_text(df)

def parse_tablepress_tables(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows_out: list[dict[str, Any]] = []
    tables = data.get("tables")
    if not isinstance(tables, list):
        return rows_out

    for table in tables:
        if not isinstance(table, dict):
            continue
        table_index = table.get("table_index")
        table_rows = table.get("rows")
        if not isinstance(table_rows, list):
            continue

        for row in table_rows:
            if not isinstance(row, dict):
                continue
            record: dict[str, Any] = {
                "table_index": table_index,
                "row_index": row.get("row_index"),
            }
            columns = row.get("columns")
            if not isinstance(columns, list):
                continue
            for column in columns:
                if not isinstance(column, dict):
                    continue
                label = column.get("header_label") or f"column_{column.get('column_index')}"
                cell = column.get("cell")
                if isinstance(cell, dict):
                    record[str(label)] = cell.get("text_stripped") or cell.get("text") or ""
            rows_out.append(record)
    return rows_out

def parse_text_elements(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows_out: list[dict[str, Any]] = []

    div_blocks = data.get("div_blocks")
    if isinstance(div_blocks, list):
        for block in div_blocks:
            if not isinstance(block, dict):
                continue
            links = block.get("links")
            if isinstance(links, list):
                for link in links:
                    if not isinstance(link, dict):
                        continue
                    rows_out.append({
                        "block_type": block.get("block_type"),
                        "text": link.get("text_stripped") or link.get("text"),
                        "href": link.get("href"),
                    })

    text_elements = data.get("text_elements")
    if isinstance(text_elements, list):
        for element in text_elements:
            if not isinstance(element, dict):
                continue
            rows_out.append({
                "element_type": element.get("element_type"),
                "text": element.get("text_stripped") or element.get("text"),
                "href": element.get("href"),
            })

    return rows_out

def _is_balad_poi(item: dict[str, Any]) -> bool:
    return (
        "name" in item
        and "token" in item
        and (
            "address" in item
            or "telephone" in item
            or isinstance(item.get("geometry"), dict)
        )
    )

def _collect_balad_pois(obj: Any, out: list[dict[str, Any]], *, source: str) -> None:
    if isinstance(obj, dict):
        if _is_balad_poi(obj):
            enriched = dict(obj)
            enriched["_source"] = source
            out.append(enriched)
            return
        for value in obj.values():
            _collect_balad_pois(value, out, source=source)
    elif isinstance(obj, list):
        if obj and all(isinstance(item, dict) and _is_balad_poi(item) for item in obj):
            for item in obj:
                enriched = dict(item)
                enriched["_source"] = source
                out.append(enriched)
            return
        for item in obj:
            _collect_balad_pois(item, out, source=source)

def _extract_balad_poi_row(item: dict[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {
        "token": item.get("token"),
        "name": item.get("name"),
        "telephone": item.get("telephone") or item.get("phone"),
        "address": item.get("address"),
        "rate": item.get("rate"),
        "category": item.get("category"),
        "source": item.get("_source"),
    }

    properties = item.get("properties")
    if isinstance(properties, dict):
        if row["name"] is None:
            row["name"] = properties.get("name")
        if row["rate"] is None:
            row["rate"] = properties.get("rate")

    geometry = item.get("geometry")
    if isinstance(geometry, dict) and geometry.get("type") == "Point":
        coordinates = geometry.get("coordinates")
        if isinstance(coordinates, list) and len(coordinates) >= 2:
            row["longitude"] = coordinates[0]
            row["latitude"] = coordinates[1]

    return {key: value for key, value in row.items() if value is not None}

def parse_balad_hub(data: dict[str, Any]) -> list[dict[str, Any]]:
    raw_items: list[dict[str, Any]] = []

    first_page = data.get("first_page")
    if isinstance(first_page, dict):
        next_data = first_page.get("__NEXT_DATA__")
        _collect_balad_pois(next_data, raw_items, source="first_page")

    pagination_responses = data.get("pagination_responses")
    if isinstance(pagination_responses, list):
        for response in pagination_responses:
            if not isinstance(response, dict):
                continue
            page = response.get("page")
            raw_payload = response.get("raw")
            _collect_balad_pois(raw_payload, raw_items, source=f"pagination_page_{page}")

    seen_tokens: set[str] = set()
    rows: list[dict[str, Any]] = []
    for item in raw_items:
        token = item.get("token")
        if isinstance(token, str) and token:
            if token in seen_tokens:
                continue
            seen_tokens.add(token)
        rows.append(_extract_balad_poi_row(item))
    return rows

def parse_unique_records(records: Any) -> list[dict[str, Any]]:
    if not isinstance(records, list):
        return []

    rows: list[dict[str, Any]] = []
    for record in records:
        if not isinstance(record, dict):
            continue

        row: dict[str, Any] = {
            "dedupe_key": record.get("dedupe_key"),
            "token": record.get("token"),
            "place_id": record.get("place_id"),
        }

        discovery_context = record.get("discovery_context")
        if isinstance(discovery_context, dict):
            row["keyword"] = discovery_context.get("keyword")
            row["city_key"] = discovery_context.get("city_key")
            row["zone_id"] = discovery_context.get("zone_id")
            row["zone_label"] = discovery_context.get("zone_label")
            row["discovery_latitude"] = discovery_context.get("latitude")
            row["discovery_longitude"] = discovery_context.get("longitude")
            row["discovery_page"] = discovery_context.get("page")
            row["transport"] = discovery_context.get("transport")

        geojson_feature = record.get("geojson_feature")
        if isinstance(geojson_feature, dict):
            properties = geojson_feature.get("properties")
            if isinstance(properties, dict):
                row["name"] = properties.get("name")
                row["rate"] = properties.get("rate")
                row["icon"] = properties.get("icon")

            geometry = geojson_feature.get("geometry")
            if isinstance(geometry, dict):
                coordinates = geometry.get("coordinates")
                if isinstance(coordinates, list) and len(coordinates) >= 2:
                    row["longitude"] = coordinates[0]
                    row["latitude"] = coordinates[1]

        rows.append(row)
    return rows

def _largest_dict_list(data: dict[str, Any]) -> list[dict[str, Any]] | None:
    best: list[dict[str, Any]] | None = None
    for value in data.values():
        if isinstance(value, list) and value and all(isinstance(item, dict) for item in value):
            if best is None or len(value) > len(best):
                best = value
    return best

def parse_json_to_records(data: Any, source_path: Path) -> list[dict[str, Any]]:
    if isinstance(data, list):
        return [
            _flatten_record(item) if isinstance(item, dict) else {"value": item}
            for item in data
        ]

    if not isinstance(data, dict):
        return [{"value": data}]

    if "unique_records" in data:
        return parse_unique_records(data.get("unique_records"))

    if "first_page" in data or "pagination_responses" in data:
        return parse_balad_hub(data)

    dealerships = data.get("dealerships")
    if isinstance(dealerships, list) and dealerships:
        return [_flatten_record(item) for item in dealerships if isinstance(item, dict)]

    table_rows = parse_tablepress_tables(data)
    if table_rows:
        return table_rows

    text_rows = parse_text_elements(data)
    if text_rows:
        return text_rows

    largest = _largest_dict_list(data)
    if largest is not None:
        return [_flatten_record(item) for item in largest]

    LOGGER.warning(
        "Unrecognized structure in %s; exporting flattened root object as a single row.",
        source_path.name,
    )
    return [_flatten_record(data)]

def load_json_file(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def load_jsonl_records(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        try:
            parsed = json.loads(stripped)
        except json.JSONDecodeError as exc:
            LOGGER.warning("Skipping %s line %d: invalid JSON (%s)", path.name, line_number, exc)
            continue
        if not isinstance(parsed, dict):
            LOGGER.warning(
                "Skipping %s line %d: expected object, got %s",
                path.name,
                line_number,
                type(parsed).__name__,
            )
            continue
        if parsed.get("status") == "failed":
            continue
        records.append(_flatten_record(parsed))
    return records


def load_source_records(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".jsonl":
        return load_jsonl_records(path)
    data = load_json_file(path)
    skip_reason = should_skip_file(path, data)
    if skip_reason:
        raise ValueError(skip_reason)
    return parse_json_to_records(data, path)


def compile_workbook(data_dir: Path, output_path: Path) -> dict[str, int]:
    data_files = discover_data_files(data_dir)
    if not data_files:
        LOGGER.warning("No JSON/JSONL files found under %s", data_dir)
        return {}

    used_sheet_names: set[str] = set()
    sheet_row_counts: dict[str, int] = {}
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for source_path in data_files:
            try:
                if SKIP_NAME_PATTERN.search(source_path.stem):
                    LOGGER.info("Skipping %s (filename contains 'test')", source_path.name)
                    continue
                if source_path.stat().st_size == 0:
                    LOGGER.info("Skipping %s (file is empty)", source_path.name)
                    continue

                records = load_source_records(source_path)
                sheet_name = make_sheet_name(source_path, used_sheet_names)
                dataframe = prepare_dataframe(records)

                dataframe.to_excel(writer, index=False, sheet_name=sheet_name)
                worksheet = writer.sheets[sheet_name]
                style_header_row(worksheet)
                autofit_columns(worksheet)

                sheet_row_counts[sheet_name] = len(dataframe)
                LOGGER.info(
                    "Sheet '%s' <= %s (%d records, %d columns)",
                    sheet_name,
                    source_path.name,
                    len(dataframe),
                    len(dataframe.columns),
                )
            except ValueError as exc:
                LOGGER.info("Skipping %s (%s)", source_path.name, exc)
            except Exception as exc:
                LOGGER.error("Failed to process %s: %s", source_path.name, exc)

    return sheet_row_counts

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Compile operational raw JSON/JSONL files into a single multi-sheet Excel workbook.",
    )
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args(argv)

    _configure_logging(args.verbose)
    data_dir = args.data_dir.resolve()
    output_path = args.output.resolve()

    if not data_dir.is_dir():
        LOGGER.error("Data directory not found: %s", data_dir)
        return 1

    sheet_counts = compile_workbook(data_dir, output_path)
    if not sheet_counts:
        LOGGER.error("No sheets were written to %s", output_path)
        return 1

    print(f"Excel file saved to: {output_path}")
    print("Transferred records per sheet:")
    for sheet_name in sorted(sheet_counts):
        print(f"  - {sheet_name}: {sheet_counts[sheet_name]} records")

    total_records = sum(sheet_counts.values())
    print(f"Total sheets: {len(sheet_counts)} | Total records: {total_records}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())