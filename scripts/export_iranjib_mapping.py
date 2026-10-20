#!/usr/bin/env python3
"""
Export the latest IranJib zero-price snapshot to Excel for car mapping.

Usage:
  docker exec -i seenad_backend python3 scripts/export_iranjib_mapping.py
  python scripts/export_iranjib_mapping.py --output data/iranjib_mapping.xlsx
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.database import SessionLocal

DEFAULT_OUTPUT = Path("data/iranjib_mapping.xlsx")
MAX_COLUMN_WIDTH = 48
PAGE_KIND_FA = {"domestic": "داخلی", "imported": "وارداتی"}
STATUS_FA = {
    "AVAILABLE": "موجود",
    "UNAVAILABLE": "ناموجود",
    "STOPPED": "توقف فروش",
    "COMING_SOON": "به زودی",
    "UNKNOWN": "نامشخص",
}
COLUMNS = [
    "شناسه ایران‌جیب (iid)",
    "کلید نگاشت",
    "صفحه",
    "سازنده",
    "نام خودرو",
    "قیمت بازار (تومان)",
    "قیمت کارخانه (تومان)",
    "وضعیت",
    "قیمت بازار خام",
    "قیمت کارخانه خام",
    "car_id",
]


def _provider_key(page_kind: str, iid: int | None, section: str, title: str) -> str:
    if iid is not None:
        return f"{page_kind}_{iid}"
    section_n = (section or "").strip().replace(" ", "_")
    title_n = (title or "").strip().replace(" ", "_")
    return f"{page_kind}_{section_n}_{title_n}"[:255]


def _autofit(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="center")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap
    for idx, column in enumerate(worksheet.columns, start=1):
        longest = 0
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            longest = max(longest, min(len(value), MAX_COLUMN_WIDTH))
            cell.alignment = Alignment(vertical="center")
        worksheet.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 12), MAX_COLUMN_WIDTH)


def _latest_batch_id(session) -> str:
    row = session.execute(
        text(
            """
            SELECT batch_id
            FROM iranjib_zero_price_scrape_snapshot
            GROUP BY batch_id
            ORDER BY MIN(created_at) DESC
            LIMIT 1
            """
        )
    ).first()
    if row is None:
        raise SystemExit("no rows in iranjib_zero_price_scrape_snapshot")
    return str(row[0])


def load_rows(*, batch_id: str | None) -> tuple[str, pd.DataFrame]:
    with SessionLocal() as session:
        chosen = batch_id or _latest_batch_id(session)
        records = session.execute(
            text(
                """
                SELECT
                    iranjib_iid,
                    page_kind,
                    section_raw,
                    title_raw,
                    market_price_toman,
                    factory_price_toman,
                    status,
                    market_price_raw,
                    factory_price_raw
                FROM iranjib_zero_price_scrape_snapshot
                WHERE batch_id = :batch_id
                ORDER BY page_kind, section_raw, title_raw, iranjib_iid
                """
            ),
            {"batch_id": chosen},
        ).mappings().all()
    if not records:
        raise SystemExit(f"no rows for batch_id={chosen}")
    rows = []
    for item in records:
        page_kind = str(item["page_kind"] or "")
        iid = item["iranjib_iid"]
        rows.append(
            {
                "شناسه ایران‌جیب (iid)": iid,
                "کلید نگاشت": _provider_key(page_kind, iid, item["section_raw"] or "", item["title_raw"] or ""),
                "صفحه": PAGE_KIND_FA.get(page_kind, page_kind),
                "سازنده": item["section_raw"] or "",
                "نام خودرو": item["title_raw"] or "",
                "قیمت بازار (تومان)": item["market_price_toman"],
                "قیمت کارخانه (تومان)": item["factory_price_toman"],
                "وضعیت": STATUS_FA.get(str(item["status"] or ""), item["status"]),
                "قیمت بازار خام": item["market_price_raw"] or "",
                "قیمت کارخانه خام": item["factory_price_raw"] or "",
                "car_id": None,
            }
        )
    return chosen, pd.DataFrame(rows, columns=COLUMNS)


def export_excel(frame: pd.DataFrame, *, output: Path, batch_id: str) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="همه", index=False)
        domestic = frame[frame["صفحه"] == "داخلی"]
        imported = frame[frame["صفحه"] == "وارداتی"]
        domestic.to_excel(writer, sheet_name="داخلی", index=False)
        imported.to_excel(writer, sheet_name="وارداتی", index=False)
        meta = pd.DataFrame(
            [
                {"فیلد": "batch_id", "مقدار": batch_id},
                {"فیلد": "generated_at", "مقدار": generated},
                {"فیلد": "rows_total", "مقدار": len(frame)},
                {"فیلد": "rows_domestic", "مقدار": len(domestic)},
                {"فیلد": "rows_imported", "مقدار": len(imported)},
                {
                    "فیلد": "note",
                    "مقدار": "کلید نگاشت پایدار: pageKind_iid (مثلاً domestic_16025). ستون car_id را با شناسه خودروی سایت پر کنید.",
                },
            ]
        )
        meta.to_excel(writer, sheet_name="اطلاعات", index=False)
        for worksheet in writer.sheets.values():
            _autofit(worksheet)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--batch-id", default=None, help="Defaults to the latest snapshot batch.")
    args = parser.parse_args(argv)
    batch_id, frame = load_rows(batch_id=args.batch_id)
    path = export_excel(frame, output=args.output, batch_id=batch_id)
    print(
        f"wrote {path} batch_id={batch_id} rows={len(frame)} "
        f"domestic={int((frame['صفحه']=='داخلی').sum())} "
        f"imported={int((frame['صفحه']=='وارداتی').sum())}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
