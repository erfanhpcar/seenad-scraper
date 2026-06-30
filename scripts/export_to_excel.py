#!/usr/bin/env python3
"""
Export scraped dealer JSON / JSONL files under data/ to Excel (.xlsx) with zero column loss.

Each *.json (top-level array) or *.jsonl (JSON Lines) file becomes its own workbook beside
the source file. All keys from the records become columns; known English keys get Persian
headers. List fields (e.g. phone_numbers, phones) are joined for Excel display; dict fields
are JSON-serialized.
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger("export-to-excel")

PERSIAN_HEADER_MAP: dict[str, str] = {
    "dealer_code": "کد نمایندگی",
    "name": "نام نمایندگی",
    "manager": "مدیریت",
    "province": "استان",
    "city": "شهر",
    "phone_numbers": "شماره‌های تماس",
    "address": "آدرس دقیق پستی",
    "agency_type": "نوع نمایندگی",
    "status": "رتبه / وضعیت",
}

DEFAULT_DATA_DIR = Path(__file__).resolve().parents[1] / "data"
MAX_COLUMN_WIDTH = 80


def _configure_logging(verbose: bool) -> None:
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        stream=sys.stderr,
    )


def discover_data_files(data_dir: Path) -> list[Path]:
    json_files = list(data_dir.glob("*.json"))
    jsonl_files = list(data_dir.glob("*.jsonl"))
    return sorted(json_files + jsonl_files)


def load_json_records(path: Path) -> list[dict[str, Any]] | None:
    """Return a list of record dicts when the file is a JSON array of objects."""
    raw = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(raw, list):
        if not raw:
            return []
        if all(isinstance(item, dict) for item in raw):
            return raw
        LOGGER.warning(
            "Skipping %s: JSON array contains non-object elements.",
            path.name,
        )
        return None
    LOGGER.warning(
        "Skipping %s: root is %s, expected a JSON array of objects.",
        path.name,
        type(raw).__name__,
    )
    return None


def load_jsonl_dataframe(path: Path) -> pd.DataFrame | None:
    """Load a JSON Lines file with pandas (one JSON object per line)."""
    if path.stat().st_size == 0:
        return pd.DataFrame()
    try:
        df = pd.read_json(path, lines=True)
    except ValueError as exc:
        LOGGER.warning("Skipping %s: invalid JSONL (%s).", path.name, exc)
        return None
    return df


def load_dataframe(path: Path) -> pd.DataFrame | None:
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        return load_jsonl_dataframe(path)
    if suffix == ".json":
        records = load_json_records(path)
        if records is None:
            return None
        return pd.DataFrame(records)
    LOGGER.warning("Skipping %s: unsupported extension %s.", path.name, path.suffix)
    return None


def _format_list_cell(values: list[Any]) -> str:
    parts: list[str] = []
    for item in values:
        if isinstance(item, (dict, list)):
            parts.append(json.dumps(item, ensure_ascii=False))
        elif item is None:
            parts.append("")
        else:
            parts.append(str(item).strip())
    return ", ".join(part for part in parts if part)


def serialize_cell(value: Any) -> Any:
    """Make nested structures Excel-safe without dropping information."""
    if value is None:
        return value
    if isinstance(value, float) and pd.isna(value):
        return value
    if isinstance(value, list):
        return _format_list_cell(value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def clean_array_fields(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for column in out.columns:
        out[column] = out[column].apply(serialize_cell)
    return out


def apply_persian_headers(df: pd.DataFrame) -> pd.DataFrame:
    rename = {key: label for key, label in PERSIAN_HEADER_MAP.items() if key in df.columns}
    return df.rename(columns=rename)


def coerce_excel_text(df: pd.DataFrame) -> pd.DataFrame:
    """Write every cell as text so Excel does not strip leading zeros or alter JSON blobs."""
    out = df.copy()
    for column in out.columns:
        out[column] = out[column].map(
            lambda value: ""
            if value is None or (isinstance(value, float) and pd.isna(value))
            else str(value)
        )
    return out


def autofit_columns(worksheet: Any, *, max_width: int = MAX_COLUMN_WIDTH) -> None:
    for column_index, column_cells in enumerate(
        worksheet.iter_cols(1, worksheet.max_column),
        start=1,
    ):
        max_length = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        width = min(max(max_length + 2, 10), max_width)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def prepare_dataframe(df: pd.DataFrame, *, source_name: str) -> pd.DataFrame:
    source_columns = list(df.columns)
    df = clean_array_fields(df)
    df = apply_persian_headers(df)
    df = coerce_excel_text(df)

    if list(df.columns) != [
        PERSIAN_HEADER_MAP.get(col, col) for col in source_columns
    ]:
        raise RuntimeError(f"Column loss detected while exporting {source_name}")

    return df


def export_data_to_excel(source_path: Path, *, output_path: Path | None = None) -> Path | None:
    df = load_dataframe(source_path)
    if df is None:
        return None

    df = prepare_dataframe(df, source_name=source_path.name)

    destination = output_path or source_path.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="dealers")
        autofit_columns(writer.sheets["dealers"])

    LOGGER.info(
        "Exported %s -> %s (%d rows, %d columns)",
        source_path.name,
        destination.name,
        len(df),
        len(df.columns),
    )
    return destination


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Export dealer JSON/JSONL files in data/ to Excel with full schema preservation.",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help="Directory containing source files (default: project data/)",
    )
    parser.add_argument(
        "--files",
        nargs="*",
        type=Path,
        help="Optional explicit .json/.jsonl paths; defaults to all in --data-dir",
    )
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args(argv)

    _configure_logging(args.verbose)

    data_dir = args.data_dir.resolve()
    if not data_dir.is_dir():
        LOGGER.error("Data directory not found: %s", data_dir)
        return 1

    if args.files:
        source_files = [path.resolve() for path in args.files]
    else:
        source_files = discover_data_files(data_dir)

    if not source_files:
        LOGGER.warning("No JSON or JSONL files found in %s", data_dir)
        return 0

    exported = 0
    skipped = 0
    for source_path in source_files:
        if not source_path.is_file():
            LOGGER.warning("File not found: %s", source_path)
            skipped += 1
            continue
        result = export_data_to_excel(source_path)
        if result is None:
            skipped += 1
        else:
            exported += 1

    LOGGER.info("Done: %d exported, %d skipped.", exported, skipped)
    return 0 if exported > 0 or skipped == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
