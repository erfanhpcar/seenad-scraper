#!/usr/bin/env python3
"""Import Top-50 zero-price benchmark watchlist from Excel into benchmark_car."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import create_engine, text


def _load_url() -> str:
    secret = Path("secrets/benchmark_database_url")
    if secret.is_file():
        return secret.read_text(encoding="utf-8").strip()
    import os

    return (
        os.getenv("BENCHMARK_DATABASE_URL", "").strip()
        or os.getenv("DATABASE_URL", "").strip()
    )


def _column_names(engine, table: str) -> set[str]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT column_name
                FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = :table
                """
            ),
            {"table": table},
        ).scalars()
    return {str(r) for r in rows}


def _read_rows(xlsx: Path) -> list[dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet_name = "50 خودرو" if "50 خودرو" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    header_row_idx = None
    header: list[Any] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        cells = list(row)
        if "Car ID" in cells:
            header_row_idx = idx
            header = cells
            break
    if header_row_idx is None:
        wb.close()
        raise SystemExit(f"No header row with 'Car ID' in sheet {sheet_name}")

    def col(name: str) -> int | None:
        try:
            return header.index(name)
        except ValueError:
            return None

    idx_rank = col("ردیف")
    idx_car = col("Car ID")
    idx_group = col("گروه")
    if idx_car is None:
        wb.close()
        raise SystemExit("Missing column: Car ID")

    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row_idx + 2, values_only=True):
        cells = list(row)
        if idx_car >= len(cells):
            continue
        car_id = cells[idx_car]
        if car_id is None or str(car_id).strip() == "":
            continue
        try:
            seenad_car_id = int(car_id)
        except (TypeError, ValueError):
            continue
        priority = None
        if idx_rank is not None and idx_rank < len(cells) and cells[idx_rank] is not None:
            try:
                priority = int(cells[idx_rank])
            except (TypeError, ValueError):
                priority = None
        group = None
        if idx_group is not None and idx_group < len(cells):
            group = cells[idx_group]
            if group is not None:
                group = str(group).strip() or None
        out.append(
            {
                "seenad_car_id": seenad_car_id,
                "priority": priority,
                "watchlist_group": group,
            }
        )
    wb.close()
    return out


def _upsert_car(
    conn,
    *,
    columns: set[str],
    seenad_car_id: int,
    priority: int | None,
    watchlist_group: str | None,
    replace: bool,
) -> None:
    fields: dict[str, Any] = {
        "seenad_car_id": seenad_car_id,
        "is_active": True,
    }
    if "priority" in columns and priority is not None:
        fields["priority"] = priority
    if "watchlist_group" in columns and watchlist_group is not None:
        fields["watchlist_group"] = watchlist_group
    if "mapping_status" in columns:
        fields["mapping_status"] = "NEEDS_MAPPING"

    keys = list(fields.keys())
    placeholders = ", ".join(f":{k}" for k in keys)
    cols_sql = ", ".join(keys)
    update_parts = [f"{k} = EXCLUDED.{k}" for k in keys if k != "seenad_car_id"]
    sql = text(
        f"""
        INSERT INTO benchmark_car ({cols_sql})
        VALUES ({placeholders})
        ON CONFLICT (seenad_car_id) DO UPDATE SET {", ".join(update_parts)}
        """
    )
    try:
        conn.execute(sql, fields)
        return
    except Exception:
        if not replace:
            existing = conn.execute(
                text("SELECT id FROM benchmark_car WHERE seenad_car_id = :cid LIMIT 1"),
                {"cid": seenad_car_id},
            ).first()
            if existing:
                sets = ", ".join(f"{k} = :{k}" for k in fields if k != "seenad_car_id")
                conn.execute(
                    text(f"UPDATE benchmark_car SET {sets} WHERE seenad_car_id = :seenad_car_id"),
                    fields,
                )
                return
        cols_sql = ", ".join(keys)
        placeholders = ", ".join(f":{k}" for k in keys)
        conn.execute(
            text(f"INSERT INTO benchmark_car ({cols_sql}) VALUES ({placeholders})"),
            fields,
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=Path("tmp/SeenAd_Top_50_Setoudeh_Priority (1).xlsx"),
        help="Watchlist Excel (sheet «50 خودرو», column Car ID).",
    )
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Deactivate benchmark_car rows not present in the file after import.",
    )
    args = parser.parse_args(argv)

    url = _load_url()
    if not url:
        print("Missing BENCHMARK_DATABASE_URL or secrets/benchmark_database_url", file=sys.stderr)
        return 1
    if not args.xlsx.is_file():
        print(f"XLSX not found: {args.xlsx}", file=sys.stderr)
        return 1

    rows = _read_rows(args.xlsx)
    if not rows:
        print("No car rows parsed from XLSX", file=sys.stderr)
        return 1

    engine = create_engine(url, pool_pre_ping=True)
    columns = _column_names(engine, "benchmark_car")
    if "seenad_car_id" not in columns:
        print("benchmark_car.seenad_car_id missing — run migration on production DB first", file=sys.stderr)
        return 2

    ids = {r["seenad_car_id"] for r in rows}
    with engine.begin() as conn:
        for row in rows:
            _upsert_car(
                conn,
                columns=columns,
                seenad_car_id=row["seenad_car_id"],
                priority=row["priority"],
                watchlist_group=row["watchlist_group"],
                replace=args.replace,
            )
        if args.replace and "is_active" in columns:
            conn.execute(
                text(
                    """
                    UPDATE benchmark_car
                    SET is_active = FALSE
                    WHERE seenad_car_id IS NOT NULL
                      AND NOT (seenad_car_id = ANY(:ids))
                    """
                ),
                {"ids": list(ids)},
            )

    with engine.connect() as conn:
        active = conn.execute(
            text("SELECT count(*) FROM benchmark_car WHERE is_active IS TRUE")
        ).scalar_one()
    print(f"imported/updated {len(rows)} watchlist cars; active benchmark_car: {active}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
